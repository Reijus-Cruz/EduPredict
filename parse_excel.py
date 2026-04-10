import openpyxl
wb = openpyxl.load_workbook('templates/FOR TESTING FINAL.xlsx', data_only=True)
print('Sheets:', wb.sheetnames)
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f'\n=== {sname} (rows={ws.max_row}, cols={ws.max_column}) ===')
    for r in range(1, min(ws.max_row+1, 50)):
        vals = []
        for c in range(1, min(ws.max_column+1, 25)):
            v = ws.cell(r, c).value
            if v is not None:
                vals.append(f'C{c}={v}')
        if vals:
            print(f'  R{r}: ' + '  '.join(vals))
